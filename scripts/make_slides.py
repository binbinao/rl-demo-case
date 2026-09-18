#!/usr/bin/env python
"""Generate the pre-sales deck "拾放智造 (PickTeach)" as a .pptx file.

Style: business-formal + tech-blue. 16:9. 11 slides with speaker notes.
Run with the project .venv (python-pptx installed):
    source .venv/bin/activate
    python scripts/make_slides.py
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

PROJECT_ROOT = "/data/robinji/rl-demo-case"
OUT_DIR = os.path.join(PROJECT_ROOT, "slides")
OUT_FILE = os.path.join(OUT_DIR, "拾放智造_售前宣讲.pptx")

# ---- palette: tech blue + business formal ----
PRIMARY = RGBColor(0x0F, 0x4C, 0x81)    # deep blue
SECONDARY = RGBColor(0x2E, 0x75, 0xB6)  # mid blue
ACCENT = RGBColor(0xED, 0x7D, 0x31)     # orange accent
LIGHT_BG = RGBColor(0xDE, 0xEB, 0xF7)   # pale blue
LIGHT_GRAY = RGBColor(0xF2, 0xF5, 0xF8) # near-white gray
DARK = RGBColor(0x33, 0x33, 0x33)       # body text
GRAY = RGBColor(0x59, 0x59, 0x59)       # secondary text
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TABLE_HEAD = RGBColor(0x1F, 0x5E, 0x9E)

FONT = "微软雅黑"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_font(run, size, color=DARK, bold=False, name=FONT, italic=False):
    """Set font incl. East-Asian typeface so Chinese renders correctly."""
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


def add_rect(slide, left, top, width, height, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp


def add_text(slide, left, top, width, height, runs, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, space_after=6, line_spacing=1.0):
    """Add a text box. `runs` is a list of paragraphs; each paragraph is a list
    of (text, size, color, bold) tuples."""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        for (text, size, color, bold) in para:
            r = p.add_run()
            r.text = text
            set_font(r, size, color, bold)
    return tb


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # blank layout


def add_header(slide, title, subtitle=None):
    """Consistent top banner: deep-blue band + white title + optional subtitle."""
    add_rect(slide, 0, 0, SLIDE_W, Inches(1.05), PRIMARY)
    add_rect(slide, 0, Inches(1.05), SLIDE_W, Inches(0.04), ACCENT)
    add_text(slide, Inches(0.55), Inches(0.16), Inches(11.5), Inches(0.6),
             [[(title, 26, WHITE, True)]])
    if subtitle:
        add_text(slide, Inches(0.57), Inches(0.68), Inches(12.0), Inches(0.35),
                 [[(subtitle, 12, RGBColor(0xD6, 0xE4, 0xF0), False)]])


def add_footer(slide, page_no):
    add_rect(slide, 0, Inches(7.16), SLIDE_W, Inches(0.34), LIGHT_GRAY)
    add_text(slide, Inches(0.55), Inches(7.20), Inches(8.0), Inches(0.25),
             [[("拾放智造 · 机械臂智能化改造方案", 9, GRAY, False)]])
    add_text(slide, Inches(12.3), Inches(7.20), Inches(0.8), Inches(0.25),
             [[(str(page_no), 10, GRAY, True)]], align=PP_ALIGN.RIGHT)


def bullets(slide, left, top, width, height, items, size=16, gap=10):
    """`items`: list of (text, level) or plain strings (level 0)."""
    runs = []
    for it in items:
        if isinstance(it, tuple):
            text, level = it
        else:
            text, level = it, 0
        bullet = "•  " if level == 0 else "–  "
        indent = "      " if level == 1 else ""
        color = DARK if level == 0 else GRAY
        bsize = size if level == 0 else size - 2
        runs.append([(indent + bullet + text, bsize, color, False)])
    add_text(slide, left, top, width, height, runs, space_after=gap, line_spacing=1.15)


# --------------------------------------------------------------------------
# Slide builders
# --------------------------------------------------------------------------

def slide_cover(prs):
    s = blank_slide(prs)
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, PRIMARY)
    add_rect(s, 0, Inches(4.9), SLIDE_W, Inches(0.06), ACCENT)
    add_text(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(1.4),
             [[("拾放智造", 60, WHITE, True)]])
    add_text(s, Inches(0.95), Inches(3.15), Inches(11.5), Inches(0.7),
             [[("小算力 × 小数据，让机械臂学会「抓放」", 24, RGBColor(0xD6, 0xE4, 0xF0), False)]])
    add_text(s, Inches(0.95), Inches(5.15), Inches(11.5), Inches(0.6),
             [[("PickTeach · 机械臂智能化改造方案", 15, RGBColor(0x9F, 0xC0, 0xDF), False)]])
    add_text(s, Inches(0.95), Inches(6.5), Inches(11.5), Inches(0.5),
             [[("对内技术复现 · 对外售前宣讲   |   SmolVLA × SO-100 案例实证", 12,
                RGBColor(0x9F, 0xC0, 0xDF), False)]])
    add_notes(s, "开场白：各位好。今天向大家汇报的是一套面向制造现场机械臂智能化改造的"
                 "方案——「拾放智造」。它要回答一个核心问题：在现有六轴/协作机器人的基础上，"
                 "能不能用很小的算力和有限的数据，就让机械臂学会「抓取-放置」这类任务，"
                 "从而降低对专业示教工程师的依赖。下面我们用一个已跑通的实际案例来说明。")


def slide_pain_points(prs):
    s = blank_slide(prs)
    add_header(s, "制造现场的「最后一公里」", "为什么机械臂智能化改造总是卡在落地")
    cards = [
        ("编程门槛高", "传统示教需专业机器人工程师\n逐点示教、反复调试，门槛高"),
        ("换品种重示教", "换产线 / 换物料就要重新示教，\n周期长、成本反复投入"),
        ("依赖现场工程师", "小批量多品种场景下，\n投入产出比不划算"),
        ("传统视觉泛化差", "光照、物料、位姿变化，\n需频繁人工调参"),
    ]
    x = Inches(0.55)
    y = Inches(1.5)
    w = Inches(2.95)
    h = Inches(4.4)
    gap = Inches(0.22)
    for i, (title, desc) in enumerate(cards):
        cx = x + i * (w + gap)
        add_rect(s, cx, y, w, h, LIGHT_BG)
        add_rect(s, cx, y, w, Inches(0.7), PRIMARY)
        add_text(s, cx + Inches(0.15), y + Inches(0.12), w - Inches(0.3), Inches(0.5),
                 [[(title, 16, WHITE, True)]])
        lines = desc.split("\n")
        runs = [[(ln, 13, DARK, False)] for ln in lines]
        add_text(s, cx + Inches(0.2), y + Inches(1.0), w - Inches(0.4), Inches(3.0),
                 runs, space_after=8, line_spacing=1.3)
    add_footer(s, 2)
    add_notes(s, "这一页讲的是「痛点」。制造现场的机械臂已经很多了，但智能化改造并没有"
                 "大规模铺开，卡在哪？四个字概括就是「落地难」：第一，传统示教要专业工程师"
                 "逐点编程，门槛高；第二，一旦换品种就要重新示教，成本反复；第三，"
                 "小批量多品种本来就是制造业常态，靠人工示教不划算；第四，传统视觉方案"
                 "对环境变化很敏感，要反复调参。这四点，就是我们方案的切入点。")


def slide_comparison(prs):
    s = blank_slide(prs)
    add_header(s, "传统示教 vs 拾放智造", "一条「由演示学习」而非「逐点编程」的路径")
    # left card
    add_rect(s, Inches(0.55), Inches(1.5), Inches(5.9), Inches(4.9), LIGHT_GRAY)
    add_rect(s, Inches(0.55), Inches(1.5), Inches(5.9), Inches(0.7), GRAY)
    add_text(s, Inches(0.75), Inches(1.62), Inches(5.5), Inches(0.5),
             [[("传统示教编程", 17, WHITE, True)]])
    bullets(s, Inches(0.85), Inches(2.5), Inches(5.4), Inches(3.6),
            ["专业工程师逐点示教，周期长",
             "换品种 = 从头再来一遍",
             "严重依赖个人经验，难以复制",
             "小批量多品种投入产出比低"])
    # right card
    add_rect(s, Inches(6.85), Inches(1.5), Inches(5.9), Inches(4.9), LIGHT_BG)
    add_rect(s, Inches(6.85), Inches(1.5), Inches(5.9), Inches(0.7), SECONDARY)
    add_text(s, Inches(7.05), Inches(1.62), Inches(5.5), Inches(0.5),
             [[("拾放智造（模仿学习微调）", 17, WHITE, True)]])
    bullets(s, Inches(7.15), Inches(2.5), Inches(5.4), Inches(3.6),
            ["演示几次给机器人看，模型自己学",
             "换品种只需补采少量数据，再微调",
             "流程可复制，不依赖个人经验",
             "单卡算力即可，适合多品种快速迭代"])
    add_footer(s, 3)
    add_notes(s, "这页做核心对比。左边是传统示教：工程师拿着示教器，一个点一个点地教机器人，"
                 "换品种就重来，靠个人经验。右边是我们的路径：模仿学习微调——"
                 "你只需要给机器人演示几次正确动作，模型从示范里学会抓放；换品种时补采少量数据"
                 "重新微调就行。注意，这里的关键词是「由示范学习」，而不是「重新编程」。")


def slide_value(prs):
    s = blank_slide(prs)
    add_header(s, "核心价值主张")
    add_text(s, Inches(0.55), Inches(1.6), Inches(12.2), Inches(1.0),
             [[("小算力 + 小数据 → 让机械臂学会抓放", 30, PRIMARY, True)]],
             align=PP_ALIGN.CENTER)
    pills = [
        ("小算力", "单张 Tesla T4（16GB）\n即可完成训练"),
        ("小数据", "几十集演示数据\n（非海量标注）"),
        ("快交付", "天级完成数据→模型\n→验证闭环"),
    ]
    w = Inches(3.7)
    h = Inches(2.4)
    gap = Inches(0.35)
    total = 3 * w + 2 * gap
    x0 = (SLIDE_W - total) / 2
    for i, (t, d) in enumerate(pills):
        cx = x0 + i * (w + gap)
        add_rect(s, cx, Inches(3.2), w, h, LIGHT_BG)
        add_rect(s, cx, Inches(3.2), w, Inches(0.18), ACCENT)
        add_text(s, cx + Inches(0.3), Inches(3.55), w - Inches(0.6), Inches(0.6),
                 [[(t, 22, PRIMARY, True)]], align=PP_ALIGN.CENTER)
        add_text(s, cx + Inches(0.3), Inches(4.25), w - Inches(0.6), Inches(1.0),
                 [[(ln, 14, DARK, False)] for ln in d.split("\n")],
                 align=PP_ALIGN.CENTER, space_after=4)
    add_footer(s, 4)
    add_notes(s, "一句话价值主张：小算力加小数据，就能让机械臂学会抓放。三个关键词——"
                 "小算力：单张 T4 卡就够，不用 GPU 集群；小数据：几十集演示数据，"
                 "不需要海量人工标注；快交付：从数据到模型到验证，天级就能闭环。"
                 "这三点分别对应客户的算力预算、数据积累、和上线节奏，是我们方案最核心的卖点。")


def slide_three_ways(prs):
    s = blank_slide(prs)
    add_header(s, "三种「教机器人」的方式", "为什么我们选择模仿学习这条更轻的路线")
    cols = [
        ("传统示教", "逐点编程", ["人工逐点示教", "换品种重来", "依赖专家经验"], GRAY),
        ("模仿学习", "从示范中学习", ["演示→模型学会", "少量数据即可", "算力要求低"], SECONDARY),
        ("强化学习", "试错+奖励", ["需大量交互试错", "奖励设计复杂", "算力门槛高"], GRAY),
    ]
    w = Inches(3.7)
    h = Inches(4.4)
    gap = Inches(0.35)
    total = 3 * w + 2 * gap
    x0 = (SLIDE_W - total) / 2
    y = Inches(1.5)
    for i, (title, sub, items, color) in enumerate(cols):
        cx = x0 + i * (w + gap)
        highlight = (i == 1)
        fill = LIGHT_BG if highlight else LIGHT_GRAY
        add_rect(s, cx, y, w, h, fill)
        add_rect(s, cx, y, w, Inches(1.15), PRIMARY if highlight else color)
        add_text(s, cx + Inches(0.25), y + Inches(0.12), w - Inches(0.5), Inches(0.5),
                 [[(title, 18, WHITE, True)]], align=PP_ALIGN.CENTER)
        add_text(s, cx + Inches(0.25), y + Inches(0.62), w - Inches(0.5), Inches(0.4),
                 [[(sub, 12, RGBColor(0xD6, 0xE4, 0xF0), False)]], align=PP_ALIGN.CENTER)
        bullet_runs = []
        for it in items:
            bullet_runs.append([("•  " + it, 13, DARK, False)])
        add_text(s, cx + Inches(0.35), y + Inches(1.4), w - Inches(0.7), Inches(2.8),
                 bullet_runs, space_after=10, line_spacing=1.2)
        if highlight:
            add_text(s, cx + Inches(0.25), y + h - Inches(0.55), w - Inches(0.5), Inches(0.4),
                     [[("✓ 本方案采用", 13, ACCENT, True)]], align=PP_ALIGN.CENTER)
    add_footer(s, 5)
    add_notes(s, "这页是技术科普，帮客户分清三种「教机器人」的方式。传统示教是逐点编程，"
                 "门槛高、换品种重来。强化学习是让机器人试错、给奖励，但需要大量交互、"
                 "奖励设计复杂、算力门槛高。模仿学习是折中且更轻的路线：给机器人演示正确动作，"
                 "它从示范里学习，数据需求小、算力低。我们选择模仿学习，正是为了匹配"
                 "制造现场「小批量、快迭代、算力有限」的约束。", )


def slide_case(prs):
    s = blank_slide(prs)
    add_header(s, "实证案例：我们做了什么", "通用模型 + 少量演示数据 → 单卡微调出抓放能力")
    steps = [
        ("通用 VLA 模型", "SmolVLA（4.5 亿参数）\n视觉-语言-动作模型"),
        ("少量演示数据", "50 集抓放演示\n约 2 万帧 · 双相机 · 6 自由度"),
        ("单卡微调", "Tesla T4（16GB）\n训练约 9 小时"),
        ("专用模型", "可加载推理的\n抓放动作模型"),
    ]
    w = Inches(2.9)
    h = Inches(3.6)
    gap = Inches(0.28)
    total = 4 * w + 3 * gap
    x0 = (SLIDE_W - total) / 2
    y = Inches(1.7)
    for i, (t, d) in enumerate(steps):
        cx = x0 + i * (w + gap)
        add_rect(s, cx, y, w, h, LIGHT_BG)
        add_rect(s, cx, y, w, Inches(0.8), PRIMARY)
        add_text(s, cx + Inches(0.2), y + Inches(0.14), w - Inches(0.4), Inches(0.55),
                 [[(t, 16, WHITE, True)]], align=PP_ALIGN.CENTER)
        add_text(s, cx + Inches(0.25), y + Inches(1.1), w - Inches(0.5), Inches(2.2),
                 [[(ln, 13, DARK, False)] for ln in d.split("\n")],
                 align=PP_ALIGN.CENTER, space_after=6, line_spacing=1.3)
        if i < 3:
            add_text(s, cx + w - Inches(0.05), y + Inches(1.4), Inches(0.4), Inches(0.5),
                     [[("→", 22, ACCENT, True)]], align=PP_ALIGN.CENTER)
    add_footer(s, 6)
    add_notes(s, "这页讲我们具体做了什么，是一条完整的流水线。起点是一个通用的大模型 SmolVLA，"
                 "4.5 亿参数，具备视觉和语言理解能力；第二步，我们只用 50 集抓放演示数据——"
                 "注意是「演示」不是海量标注，大约两万帧、双相机、6 自由度；第三步，在单张 T4 "
                 "16G 显卡上微调约 9 小时；最后得到一个可加载推理的专用抓放模型。"
                 "整条链路没有用到 GPU 集群，成本可控。")


def slide_results(prs):
    s = blank_slide(prs)
    add_header(s, "效果数据", "离线动作预测对比：微调后误差大幅下降")
    # big number
    add_rect(s, Inches(0.55), Inches(1.4), Inches(4.0), Inches(2.6), PRIMARY)
    add_text(s, Inches(0.75), Inches(1.7), Inches(3.6), Inches(1.1),
             [[("-89.3%", 44, WHITE, True)]], align=PP_ALIGN.CENTER)
    add_text(s, Inches(0.75), Inches(2.9), Inches(3.6), Inches(0.7),
             [[("动作预测误差下降", 15, RGBColor(0xD6, 0xE4, 0xF0), False)]],
             align=PP_ALIGN.CENTER)
    # table
    rows = [
        ("关节", "base MSE", "微调后 MSE", "提升"),
        ("shoulder_pan", "604.11", "15.82", "+97.4%"),
        ("shoulder_lift", "409.03", "75.52", "+81.5%"),
        ("elbow_flex", "183.10", "33.09", "+81.9%"),
        ("wrist_flex", "202.49", "23.49", "+88.4%"),
        ("wrist_roll", "98.07", "15.57", "+84.1%"),
        ("gripper", "92.69", "6.83", "+92.6%"),
        ("总体", "264.91", "28.39", "+89.3%"),
    ]
    tbl_x = Inches(4.9)
    tbl_y = Inches(1.4)
    tbl_w = Inches(7.85)
    rows_n = len(rows)
    cols_n = 4
    gtable = s.shapes.add_table(rows_n, cols_n, tbl_x, tbl_y, tbl_w,
                                Inches(0.42) * rows_n).table
    gtable.columns[0].width = Inches(2.3)
    gtable.columns[1].width = Inches(1.8)
    gtable.columns[2].width = Inches(1.85)
    gtable.columns[3].width = Inches(1.9)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = gtable.cell(ri, ci)
            cell.text = val
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER
            for run in para.runs:
                if ri == 0:
                    set_font(run, 12, WHITE, True)
                elif ri == rows_n - 1:
                    set_font(run, 12, PRIMARY, True)
                else:
                    set_font(run, 12, DARK, False)
            if ri == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = TABLE_HEAD
            elif ri == rows_n - 1:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_BG
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_GRAY if ri % 2 else WHITE
    add_text(s, Inches(0.55), Inches(4.35), Inches(12.2), Inches(2.4),
             [[("说明：离线评估（留出集 ep45–49，共 1877 帧），指标为预测动作与真值动作的逐关节 MSE，", 11, GRAY, False)],
              [("数值越低越准。微调后 6 个关节全部提升（+81.5% ~ +97.4%）。", 11, GRAY, False)]],
             space_after=4)
    add_footer(s, 7)
    add_notes(s, "这一页是硬数据，最有说服力。我们做一个「离线动作预测对比」：把模型预测的动作"
                 "和数据集里的真值动作逐关节比对，算均方误差，误差越小越准。结果：微调后的模型"
                 "整体误差下降了 89.3%，从 264.91 降到 28.39；6 个关节全部提升，"
                 "最好的一档提升了 97.4%。需要说明这是离线验证、不是真机跑，"
                 "但它已经能证明：模型确实学到了这个抓放任务的动作规律。")


def slide_cost(prs):
    s = blank_slide(prs)
    add_header(s, "成本账", "算力 + 数据 + 人力 + 工期（口径可替换为贵司真实值）")
    rows = [
        ("维度", "传统示教编程", "拾放智造（本案例推算）"),
        ("算力要求", "无需训练（纯人工编程）", "单张 T4，约 9 小时"),
        ("数据投入", "—", "50 集演示（约 1–2 人·天采集）"),
        ("单工位总投入", "约 5–15 人·天", "约 4–6 人·天"),
        ("换品种成本", "重新示教 3–10 人·天 / 次", "补采少量数据 + 再微调（天级）"),
        ("专家依赖", "高（依赖示教经验）", "中（数据采集即可上手）"),
    ]
    gtable = s.shapes.add_table(len(rows), 3, Inches(0.55), Inches(1.5),
                                Inches(12.2), Inches(0.55) * len(rows)).table
    gtable.columns[0].width = Inches(2.0)
    gtable.columns[1].width = Inches(5.1)
    gtable.columns[2].width = Inches(5.1)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = gtable.cell(ri, ci)
            cell.text = val
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.LEFT
            for run in para.runs:
                if ri == 0:
                    set_font(run, 13, WHITE, True)
                elif ci == 0:
                    set_font(run, 12, PRIMARY, True)
                else:
                    set_font(run, 12, DARK, False)
            if ri == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = TABLE_HEAD
            elif ri % 2:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_GRAY
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE
    add_text(s, Inches(0.55), Inches(5.15), Inches(12.2), Inches(1.4),
             [[("注：本案例为演示规模，实际产线需按工位数量、品种复杂度、数据采集条件重新测算；", 11, GRAY, False)],
              [("「拾放智造」列的 1–2 人·天为数据采集假设，训练为夜间无人值守（约 0.5 人·天值守），部署联调 2–3 人·天。", 11, GRAY, False)]],
             space_after=4)
    add_footer(s, 8)
    add_notes(s, "这页算成本账，口径是「算力 + 数据 + 人力 + 工期」。对比传统示教："
                 "一个复杂抓放工位，传统示教大约要 5 到 15 人天；我们这套方案推算下来约 4 到 6 人天，"
                 "而且算力只要一张 T4 卡。更关键的是「换品种」——传统要重新示教、每次 3 到 10 人天；"
                 "我们只需补采少量数据再微调，天级就能完成。表里的数字是演示规模推算的合理假设，"
                 "落地时请替换成贵司真实工位和人力口径。", )


def slide_roadmap(prs):
    s = blank_slide(prs)
    add_header(s, "落地路径", "从 demo 到客户产线：四步闭环")
    steps = [
        ("采数据", "在目标工位录制\n抓放演示（几十集）"),
        ("微调", "通用模型 + 现场数据\n单卡微调"),
        ("部署", "加载模型到推理端\n接机械臂控制"),
        ("迭代", "根据上线反馈\n补数据再微调"),
    ]
    w = Inches(2.9)
    h = Inches(3.4)
    gap = Inches(0.28)
    total = 4 * w + 3 * gap
    x0 = (SLIDE_W - total) / 2
    y = Inches(1.9)
    for i, (t, d) in enumerate(steps):
        cx = x0 + i * (w + gap)
        add_rect(s, cx, y, w, h, LIGHT_BG)
        add_rect(s, cx, y, w, w, PRIMARY if i % 2 == 0 else SECONDARY)
        add_text(s, cx, y + w / 2 - Inches(0.3), w, Inches(0.6),
                 [[(str(i + 1), 32, WHITE, True)]], align=PP_ALIGN.CENTER)
        add_text(s, cx + Inches(0.2), y + w + Inches(0.2), w - Inches(0.4), Inches(0.5),
                 [[(t, 17, PRIMARY, True)]], align=PP_ALIGN.CENTER)
        add_text(s, cx + Inches(0.25), y + w + Inches(0.75), w - Inches(0.5), Inches(1.5),
                 [[(ln, 12, DARK, False)] for ln in d.split("\n")],
                 align=PP_ALIGN.CENTER, space_after=4, line_spacing=1.3)
        if i < 3:
            add_text(s, cx + w - Inches(0.05), y + Inches(1.1), Inches(0.4), Inches(0.5),
                     [[("→", 22, ACCENT, True)]], align=PP_ALIGN.CENTER)
    add_footer(s, 9)
    add_notes(s, "落地分四步，形成一个可持续迭代的闭环。第一步采数据：在客户目标工位录制抓放演示，"
                 "几十集即可；第二步微调：用通用模型加载现场数据，单卡微调；第三步部署："
                 "把模型加载到推理端，接上机械臂控制；第四步迭代：根据上线反馈补数据再微调，"
                 "越用越准。这套流程可复制到不同工位、不同品种。")


def slide_limits(prs):
    s = blank_slide(prs)
    add_header(s, "边界与诚实声明", "把当前能力边界讲清楚，是对客户负责")
    bullets(s, Inches(0.7), Inches(1.6), Inches(12.0), Inches(4.6),
            ["当前成果为「离线动作预测验证」，尚未接真机 rollout 闭环",
             "演示场景为固定抓放任务，多品种 / 复杂工况需扩展数据",
             "本案例使用演示规模数据（50 集），产线落地需采集现场真实数据",
             "模型泛化依赖数据覆盖度：新物料、新位姿需补充演示",
             "上线前需在真实产线做二次验证与安全评估"],
            size=17, gap=14)
    add_footer(s, 10)
    add_notes(s, "这页我们主动讲边界，这是技术方案该有的诚实。第一，目前是离线验证，"
                 "证明的是「预测动作接近真值」，还没接真机跑闭环；第二，演示场景是固定抓放，"
                 "多品种复杂工况要扩数据；第三，50 集是演示规模，产线要用现场真实数据；"
                 "第四，模型泛化依赖数据覆盖；第五，上线前必须真机二次验证和安全评估。"
                 "把这些讲清楚，既是对客户负责，也避免过度承诺。")


def slide_next(prs):
    s = blank_slide(prs)
    add_header(s, "下一步", "用一个小型 POC 验证产线可行性")
    steps = [
        ("选定 1 个典型工位", "抓放任务清晰、数据易采集"),
        ("采集现场演示数据", "约定集数与采集规范"),
        ("微调 + 离线验证", "输出 MSE 评估报告"),
        ("真机部署 + 联合验收", "形成可复制的改造方案"),
    ]
    y = Inches(1.5)
    for i, (t, d) in enumerate(steps):
        add_rect(s, Inches(0.7), y + i * Inches(1.15), Inches(0.5), Inches(0.5), PRIMARY)
        add_text(s, Inches(0.7), y + i * Inches(1.15), Inches(0.5), Inches(0.5),
                 [[(str(i + 1), 18, WHITE, True)]], align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, Inches(1.5), y + i * Inches(1.15) - Inches(0.08), Inches(6.0), Inches(0.5),
                 [[(t, 17, DARK, True)]])
        add_text(s, Inches(7.6), y + i * Inches(1.15) - Inches(0.05), Inches(5.0), Inches(0.5),
                 [[(d, 13, GRAY, False)]])
    add_rect(s, Inches(0.7), Inches(6.4), Inches(12.0), Inches(0.02), PRIMARY)
    add_text(s, Inches(0.7), Inches(6.55), Inches(12.0), Inches(0.5),
             [[("诚邀贵司提供 1 个真实抓放工位，共同完成一次端到端 POC 验证。", 16, PRIMARY, True)]])
    add_footer(s, 11)
    add_notes(s, "最后一页是行动号召。建议从小型 POC 开始：选一个抓放任务清晰的典型工位，"
                 "采集现场演示数据，做微调和离线验证，最后真机部署联合验收，形成可复制的方案。"
                 "我们诚邀贵司提供真实工位，一起跑通一次端到端验证，用真实数据说话。")


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    for fn in [slide_cover, slide_pain_points, slide_comparison, slide_value,
               slide_three_ways, slide_case, slide_results, slide_cost,
               slide_roadmap, slide_limits, slide_next]:
        fn(prs)
    os.makedirs(OUT_DIR, exist_ok=True)
    prs.save(OUT_FILE)
    print(f"deck saved: {OUT_FILE} ({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    main()
